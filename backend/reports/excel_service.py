from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.drawing.image import Image as XLImage
import matplotlib.pyplot as plt
from PIL import Image
import base64


def _style_header(row):
    bold = Font(bold=True, color="FFFFFFFF")
    fill = PatternFill(fill_type="solid", start_color="FF6246EA", end_color="FF6246EA")
    for cell in row:
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))


def _add_table(ws, headers, rows):
    ws.append(headers)
    _style_header(ws[1])
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    # auto width
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                v = str(cell.value or "")
            except Exception:
                v = ""
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)


def _render_chart_image(chart_type, title, labels, values, colors=None, figsize=(8, 4)):
    plt.close("all")
    fig, ax = plt.subplots(figsize=figsize)
    try:
        if chart_type == "line":
            ax.plot(labels, values, color=colors[0] if colors else "#6246ea", marker="o")
        elif chart_type == "bar":
            ax.bar(labels, values, color=colors or ["#6246ea"])
        elif chart_type == "doughnut":
            wedges, _ = ax.pie(values, labels=labels, colors=colors, wedgeprops=dict(width=0.4))
        elif chart_type == "pie":
            ax.pie(values, labels=labels, colors=colors)
        ax.set_title(title, fontsize=12)
        ax.set_facecolor("#0b1020")
        fig.patch.set_facecolor("#0b1020")
        ax.title.set_color("#ffffff")
        for spine in ax.spines.values():
            spine.set_visible(False)
        # Rotate labels if many
        if len(labels) > 6:
            plt.setp(ax.get_xticklabels(), rotation=45, ha="right")
    except Exception:
        pass

    buf = BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png", dpi=150, facecolor=fig.get_facecolor())
    buf.seek(0)
    return buf


def generate_excel(payload):
    wb = Workbook()
    wb.remove(wb.active)

    # KPIs sheet
    ws_kpi = wb.create_sheet("KPIs")
    headers = ["Metric", "Value"]
    rows = []
    kpis = payload.get("kpis", {})
    rows.append({"Metric": "Revenue", "Value": kpis.get("revenue")})
    rows.append({"Metric": "Orders", "Value": kpis.get("orders")})
    rows.append({"Metric": "Customers", "Value": kpis.get("customers")})
    rows.append({"Metric": "AOV", "Value": kpis.get("aov")})
    _add_table(ws_kpi, headers, rows)

    # Forecast series
    f = payload.get("forecasting", {})
    if f.get("series"):
        labels = f["series"].get("labels", [])
        predicted = f["series"].get("predicted", [])
        ws = wb.create_sheet("Forecast Series")
        _add_table(ws, ["Label", "Value"], [{"Label": l, "Value": (predicted[i] if i < len(predicted) else None)} for i, l in enumerate(labels)])
        imgbuf = _render_chart_image("line", "Sales Forecast", labels, [float(v or 0) for v in predicted])
        img = XLImage(imgbuf)
        ws.add_image(img, "D2")

    # Inventory
    inv = payload.get("inventory", {})
    ws_inv = wb.create_sheet("Inventory")
    inv_rows = [{"Metric": k, "Value": v} for k, v in inv.items()]
    _add_table(ws_inv, ["Metric", "Value"], inv_rows)
    inv_labels = ["High Risk", "Medium Risk", "Low Risk"]
    inv_vals = [inv.get("high_risk", 0), inv.get("medium_risk", 0), inv.get("low_risk", 0)]
    imgbuf = _render_chart_image("bar", "Inventory Risk", inv_labels, inv_vals, colors=["#ef4444", "#f7c945", "#22c55e"])
    ws_inv.add_image(XLImage(imgbuf), "D2")

    # Pricing
    pricing = payload.get("pricing", {})
    ws_pr = wb.create_sheet("Pricing")
    pr_rows = [{"Metric": k, "Value": v} for k, v in pricing.items()]
    _add_table(ws_pr, ["Metric", "Value"], pr_rows)
    acc = pricing.get("acceptance_rate") or 0
    imgbuf = _render_chart_image("doughnut", "Pricing Acceptance", ["Accepted", "Pending"], [acc, 100 - acc], colors=["#6246ea", "#e0d7ff"]) 
    ws_pr.add_image(XLImage(imgbuf), "D2")

    # Recommendations
    reco = payload.get("recommendations", {})
    ws_rec = wb.create_sheet("Recommendations")
    rec_rows = [{"Metric": k, "Value": v} for k, v in reco.items()]
    _add_table(ws_rec, ["Metric", "Value"], rec_rows)
    reco_labels = ["Active Products", "Interactions", "Orders"]
    reco_vals = [reco.get("active_products", 0), reco.get("total_interactions", 0), reco.get("total_orders", 0)]
    imgbuf = _render_chart_image("bar", "Recommendations", reco_labels, reco_vals)
    ws_rec.add_image(XLImage(imgbuf), "D2")

    # Segmentation
    seg = payload.get("segmentation", {})
    seg_rows = [{"Segment": k, "Count": v} for k, v in (seg.get("segments") or {}).items()]
    ws_seg = wb.create_sheet("Segmentation")
    _add_table(ws_seg, ["Segment", "Count"], seg_rows)
    seg_labels = [r["Segment"] for r in seg_rows]
    seg_vals = [r["Count"] for r in seg_rows]
    if seg_labels:
        imgbuf = _render_chart_image("pie", "Customer Segmentation", seg_labels, seg_vals)
        ws_seg.add_image(XLImage(imgbuf), "D2")

    # Seasonal
    s = payload.get("seasonal", {})
    ws_se = wb.create_sheet("Seasonal")
    if s.get("series"):
        labels = s["series"].get("labels", [])
        vals = s["series"].get("predicted", [])
        _add_table(ws_se, ["Label", "Value"], [{"Label": l, "Value": (vals[i] if i < len(vals) else None)} for i, l in enumerate(labels)])
        imgbuf = _render_chart_image("bar", "Seasonal Trends", labels, [float(v or 0) for v in vals])
        ws_se.add_image(XLImage(imgbuf), "D2")

    # Final workbook -> bytes
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out
